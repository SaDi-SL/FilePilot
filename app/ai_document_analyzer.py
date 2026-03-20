"""
ai_document_analyzer.py — AI-powered document analysis for FilePilot.

Reads file content, extracts key information:
- Document type and category
- Important dates (expiry, due dates, deadlines)
- Key entities (names, amounts, locations)
- Smart folder suggestion
- Actionable tips for the user

Supports: PDF, DOCX, TXT, images (OCR), XLSX
"""
from __future__ import annotations

import json
import logging
import threading
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Callable

logger = logging.getLogger(__name__)


# ── Result dataclasses ────────────────────────────────────────────────────────

@dataclass
class ExtractedDate:
    """An important date found in the document."""
    label:       str          # "Contract expiry", "Payment due", etc.
    date:        str          # ISO format: "2025-03-15"
    description: str          # Human-readable context
    remind_days_before: int = 7  # Suggest reminder X days before


@dataclass
class DocumentAnalysis:
    """Full analysis result for a document."""
    filename:        str
    doc_type:        str                    # "contract", "invoice", "receipt", etc.
    category:        str                    # Suggested FilePilot category
    smart_folder:    str                    # Suggested subfolder path
    summary:         str                    # 1-2 sentence summary
    key_dates:       list[ExtractedDate] = field(default_factory=list)
    entities:        dict[str, str]     = field(default_factory=dict)  # name, amount, etc.
    tips:            list[str]          = field(default_factory=list)
    confidence:      float              = 0.8
    provider:        str                = "ollama"
    error:           str | None         = None

    @property
    def ok(self) -> bool:
        return self.error is None and bool(self.category)

    @property
    def has_dates(self) -> bool:
        return len(self.key_dates) > 0


# ── Content extractors ────────────────────────────────────────────────────────

def _extract_text(file_path: Path, max_chars: int = 3000) -> str:
    """Extract text from any supported file type."""
    suffix = file_path.suffix.lower()

    try:
        # Plain text files
        if suffix in (".txt", ".md", ".csv", ".log"):
            return file_path.read_text(encoding="utf-8", errors="ignore")[:max_chars]

        # PDF
        if suffix == ".pdf":
            return _extract_pdf(file_path, max_chars)

        # Word documents
        if suffix in (".docx", ".doc"):
            return _extract_docx(file_path, max_chars)

        # Excel
        if suffix in (".xlsx", ".xls"):
            return _extract_excel(file_path, max_chars)

        # Images — OCR
        if suffix in (".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".webp"):
            return _extract_image_ocr(file_path, max_chars)

        # Fallback — filename only
        return f"[Cannot read content of {suffix} file]"

    except Exception as e:
        logger.warning(f"Content extraction failed for {file_path.name}: {e}")
        return f"[Extraction error: {e}]"


def _extract_pdf(file_path: Path, max_chars: int) -> str:
    try:
        import pypdf
        reader = pypdf.PdfReader(str(file_path))
        text = ""
        for page in reader.pages[:5]:  # First 5 pages
            text += page.extract_text() or ""
            if len(text) >= max_chars:
                break
        return text[:max_chars]
    except ImportError:
        return "[pypdf not installed — pip install pypdf]"


def _extract_docx(file_path: Path, max_chars: int) -> str:
    try:
        import docx
        doc = docx.Document(str(file_path))
        text = "\n".join(p.text for p in doc.paragraphs)
        return text[:max_chars]
    except ImportError:
        return "[python-docx not installed — pip install python-docx]"


def _extract_excel(file_path: Path, max_chars: int) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
        ws = wb.active
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 50:
                break
            row_text = " | ".join(str(c) for c in row if c is not None)
            if row_text.strip():
                rows.append(row_text)
        return "\n".join(rows)[:max_chars]
    except ImportError:
        return "[openpyxl not installed — pip install openpyxl]"


def _extract_image_ocr(file_path: Path, max_chars: int) -> str:
    try:
        import pytesseract
        from PIL import Image
        img = Image.open(str(file_path))
        text = pytesseract.image_to_string(img)
        return text[:max_chars]
    except ImportError:
        return "[pytesseract not installed — pip install pytesseract]"
    except Exception as e:
        return f"[OCR failed: {e}]"


# ── AI prompt ─────────────────────────────────────────────────────────────────

def _build_analysis_prompt(filename: str, content: str, categories: list[str]) -> str:
    today = datetime.now().strftime("%Y-%m-%d")
    cats = ", ".join(f'"{c}"' for c in categories) if categories else "invoices, contracts, documents, reports, images, resumes, others"

    return f"""You are a smart document analyzer. Read the document carefully and identify what it ACTUALLY is.

Today's date: {today}
Filename: {filename}
Available categories: {cats}

Document content:
---
{content[:2500]}
---

IMPORTANT: Identify the document type from its ACTUAL content, not from assumptions.
Common types: resume/CV, invoice, contract, receipt, report, letter, certificate, medical record, ID, form, manual, article, spreadsheet, image, other

Return ONLY a valid JSON object:
{{
  "doc_type": "actual document type (e.g. resume, invoice, contract, medical, certificate, report, letter, manual, form, other)",
  "category": "most fitting category from the available list",
  "smart_folder": "Category/Subcategory/Year (based on actual content)",
  "summary": "1-2 sentences describing what this document ACTUALLY is and contains",
  "key_dates": [
    {{
      "label": "Descriptive label for this date",
      "date": "YYYY-MM-DD",
      "description": "Why this date matters",
      "remind_days_before": 7
    }}
  ],
  "entities": {{
    "key relevant fields based on doc type": "value"
  }},
  "tips": [
    "Practical tip specific to THIS document type and content"
  ],
  "confidence": 0.9
}}

Strong indicators by type:
- RESUME/CV: contains work experience, education, skills, personal info, job titles, references → doc_type = "resume"
- INVOICE: contains invoice number, billing, payment terms, line items → doc_type = "invoice"  
- CONTRACT: contains parties, terms, signatures, legal language → doc_type = "contract"
- CERTIFICATE: contains awarded to, completion, issued by → doc_type = "certificate"
- MEDICAL: contains patient, diagnosis, prescription, doctor → doc_type = "medical"
- REPORT: contains analysis, findings, conclusions, charts → doc_type = "report"
- LETTER: contains Dear/To, formal greeting, single topic, short → doc_type = "letter"

Filename hints:
- "CV" or "Resume" in filename → almost certainly a resume
- "Invoice" or "INV" → invoice
- "Contract" or "Agreement" → contract

For RESUME/CV specifically:
- entities: name, current_role, email, phone, location
- key_dates: [] (empty — no reminders needed)
- smart_folder: "Resumes/PersonName" or "Career/CVs"
- tips: "Keep this CV updated", "Store in Career folder for easy access"

Rules:
- NEVER assume invoice without seeing invoice numbers or billing info
- Filename containing "CV" = resume, always
- Only add key_dates for genuinely important future deadlines
- Return ONLY the JSON, no extra text"""



# ── Main analyzer ─────────────────────────────────────────────────────────────

class AIDocumentAnalyzer:
    """
    Analyzes document content using AI.
    Extracts type, dates, entities, and actionable tips.
    """

    def __init__(self, ai_classifier=None) -> None:
        self._ai = ai_classifier  # AIClassifier instance

    def _get_ai(self):
        if self._ai:
            return self._ai
        # Lazy import to avoid circular imports
        from app.ai_classifier import AIClassifier
        self._ai = AIClassifier()
        return self._ai

    def analyze(
        self,
        file_path: Path,
        categories: list[str] | None = None,
    ) -> DocumentAnalysis:
        """
        Analyze a document synchronously.
        Returns DocumentAnalysis with all extracted information.
        """
        file_path = Path(file_path)
        filename  = file_path.name

        # Extract text content
        content = _extract_text(file_path)

        # Build and send prompt
        ai = self._get_ai()
        provider = ai.get_active_provider()

        if provider == "none":
            return DocumentAnalysis(
                filename=filename,
                doc_type="unknown",
                category="others",
                smart_folder="others",
                summary="AI not available",
                error="No AI provider running",
            )

        try:
            prompt   = _build_analysis_prompt(filename, content, categories or [])
            backend  = ai._claude if provider == "claude" else ai._ollama
            # Document analysis needs more time than simple classification
            orig_timeout = getattr(backend, "_timeout", None)
            try:
                backend._timeout = 120  # 2 minutes for full document analysis
            except Exception:
                pass
            response = backend.chat(prompt)
            if orig_timeout is not None:
                try:
                    backend._timeout = orig_timeout
                except Exception:
                    pass
            return self._parse_response(filename, response, provider)

        except Exception as e:
            logger.error(f"Analysis failed for {filename}: {e}")
            return DocumentAnalysis(
                filename=filename,
                doc_type="unknown",
                category="others",
                smart_folder="others",
                summary=f"Analysis failed: {e}",
                error=str(e),
            )

    def analyze_async(
        self,
        file_path: Path,
        categories: list[str] | None,
        on_done: Callable[[DocumentAnalysis], None],
    ) -> None:
        """Analyze in background thread. Calls on_done(analysis) when complete."""
        def _run():
            result = self.analyze(file_path, categories)
            on_done(result)
        threading.Thread(target=_run, daemon=True).start()

    def _parse_response(
        self, filename: str, response: str, provider: str
    ) -> DocumentAnalysis:
        """Parse JSON response into DocumentAnalysis."""
        try:
            start = response.find("{")
            end   = response.rfind("}") + 1
            if start == -1 or end == 0:
                raise ValueError("No JSON in response")

            data = json.loads(response[start:end])

            # Parse key dates
            key_dates = []
            for d in data.get("key_dates", []):
                try:
                    key_dates.append(ExtractedDate(
                        label=d.get("label", "Important date"),
                        date=d.get("date", ""),
                        description=d.get("description", ""),
                        remind_days_before=int(d.get("remind_days_before", 7)),
                    ))
                except Exception:
                    pass

            return DocumentAnalysis(
                filename=filename,
                doc_type=data.get("doc_type", "other"),
                category=data.get("category", "others"),
                smart_folder=data.get("smart_folder", "others"),
                summary=data.get("summary", ""),
                key_dates=key_dates,
                entities=data.get("entities", {}),
                tips=data.get("tips", []),
                confidence=float(data.get("confidence", 0.8)),
                provider=provider,
            )

        except Exception as e:
            logger.warning(f"Parse error: {e}\nResponse: {response[:300]}")
            return DocumentAnalysis(
                filename=filename,
                doc_type="unknown",
                category="others",
                smart_folder="others",
                summary="Could not parse AI response",
                error=str(e),
            )